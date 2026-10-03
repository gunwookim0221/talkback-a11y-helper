import json
from copy import deepcopy
from types import SimpleNamespace

import pytest

from tb_runner import global_navigation as nav
from tb_runner.scroll_reliability import viewport
from tb_runner.traversal_reliability import instance_id, termination_status


NAMES = ["Home", "Devices", "Life", "Routines", "Menu"]
CONFIG = dict(scenario_id="global_nav_main", scenario_type="global_nav", max_steps=10,
              global_nav=dict(labels=NAMES, resource_ids=["nav:"+n.lower() for n in NAMES]))


def nodes(selected="home", count=5):
    result = [dict(text=name, contentDescription=name, role="tab", viewIdResourceName="nav:"+name.lower(),
                   boundsInScreen=f"{index*100},900,{index*100+90},1000", selected=name.lower()==selected,
                   clickable=True, focusable=True, visibleToUser=True) for index,name in enumerate(NAMES[:count])]
    result += [dict(text="Home card", viewIdResourceName="body", boundsInScreen="0,100,300,400", clickable=True)]
    return result


def observation(selected="home", count=5):
    n=nodes(selected,count)
    return dict(items=nav.discover(n,CONFIG),nodes=n,viewport=viewport(n,"global_nav_main"))


class Client:
    def __init__(self,selected="home",count=5,activate=True,change=True):
        self.selected=selected;self.count=count;self.activate=activate;self.change=change;self.calls=[]
    def dump_tree(self,**kwargs):
        return nodes(self.selected,self.count)
    def touch_point(self,dev,x,y):
        self.calls.append((x,y))
        if self.activate and self.change:
            self.selected=NAMES[x//100].lower()
        return self.activate
    def smart_next(self,**kwargs):
        raise AssertionError("Global Nav must not enter body SMART_NEXT")


@pytest.fixture(autouse=True)
def quiet(monkeypatch):
    monkeypatch.setattr(nav.time,"sleep",lambda _:None)
    monkeypatch.setattr(nav,"log",lambda _:None)


def run(client,tmp_path,config=None):
    rows=[]
    nav.collect(client,"d",config or CONFIG,rows,str(tmp_path/"result.xlsx"),str(tmp_path))
    return rows,client.last_main_traversal_summary


def test_five_items_discovered():
    assert len(nav.discover(nodes(),CONFIG))==5


def test_body_candidate_is_excluded():
    assert all(i["view_id"]!="body" for i in nav.discover(nodes(),CONFIG))


def test_current_selected_state():
    assert nav.current(nav.discover(nodes("devices"),CONFIG))["logical_name"]=="devices"


def test_label_alone_cannot_prove_current():
    assert nav.current(nav.discover(nodes(None),CONFIG)) is None


def test_ambiguous_selected_state_is_not_current():
    n=nodes();n[1]["selected"]=True
    assert nav.current(nav.discover(n,CONFIG)) is None


def test_same_label_different_bounds_is_distinct():
    n=nodes();other=deepcopy(n[0]);other["boundsInScreen"]="600,900,690,1000";n.append(other)
    home=[i for i in nav.discover(n,CONFIG) if i["logical_name"]=="home"]
    assert len(home)==2 and home[0]["instance_id"]!=home[1]["instance_id"]


def test_redump_identity_is_stable():
    assert [i["instance_id"] for i in nav.discover(nodes(),CONFIG)]==[i["instance_id"] for i in nav.discover(nodes("devices"),CONFIG)]


def test_identity_reuses_phase0a():
    item=nav.discover(nodes(),CONFIG)[0]
    assert item["instance_id"]==instance_id(dict(item,scenario_id=CONFIG["scenario_id"],label=item["semantic_label"]))


def test_container_path_identity_fallback():
    a=dict(scenario_id="g",label="Home",container_id="a",stable_node_path="1")
    assert instance_id(a)!=instance_id(dict(a,container_id="b"))


def test_unseen_preferred():
    assert nav.next_unseen(nav.discover(nodes(),CONFIG),{"home"},{})["logical_name"]=="devices"


def test_current_not_reactivated(tmp_path):
    c=Client();rows,s=run(c,tmp_path)
    assert len(c.calls)==4 and rows[0]["nav_state"]=="CURRENT"
    assert rows[0]["nav_activation_attempted"] is False


def test_selected_change_verifies():
    before,after=observation(),observation("devices")
    assert nav.verify_transition(before,after,before["items"][1],True)["destination_verified"]


def test_click_success_without_selected_change_not_verified():
    before=observation()
    t=nav.verify_transition(before,before,before["items"][1],True)
    assert t["result"]=="ACTIVATED_NOT_VERIFIED" and not t["destination_verified"]


def test_destination_marker_strengthens_evidence():
    before,after=observation(None),observation(None)
    after["nodes"].append(dict(text="Devices root marker"))
    assert nav.verify_transition(before,after,before["items"][1],True,["Devices root marker"])["destination_verified"]


def test_generic_body_change_cannot_prove_target():
    before,after=observation(),observation()
    after["viewport"]=viewport([dict(viewIdResourceName="new",boundsInScreen="0,100,300,400")],"g")
    t=nav.verify_transition(before,after,before["items"][1],True)
    assert t["screen_fingerprint_changed"] and not t["destination_verified"]


def test_wrong_selected_destination_rejects_marker():
    before,after=observation(),observation("life")
    after["nodes"].append(dict(text="Devices marker"))
    assert not nav.verify_transition(before,after,before["items"][1],True,["Devices marker"])["destination_verified"]


def test_activation_failure(tmp_path):
    c=Client(activate=False);_,s=run(c,tmp_path)
    assert s["nav_activation_failures"]==8 and s["nav_items_verified"]==1
    assert s["termination_status"]=="INCOMPLETE_GLOBAL_NAV"


def test_focused_not_activated(tmp_path):
    c=Client(activate=False)
    c.get_focus=lambda **kwargs:dict(accessibilityFocused=True,boundsInScreen=f"{c.calls[-1][0]-45},900,{c.calls[-1][0]+45},1000")
    rows,_=run(c,tmp_path)
    assert rows[1]["nav_state"]=="FOCUSED_NOT_ACTIVATED"


def test_all_five_only_completion(tmp_path):
    _,s=run(Client(),tmp_path)
    assert s["nav_items_verified"]==5 and s["termination_status"]=="COMPLETED"


def test_four_of_five_incomplete(tmp_path):
    _,s=run(Client(count=4),tmp_path)
    assert s["nav_items_verified"]==4 and s["termination_status"]=="INCOMPLETE_GLOBAL_NAV"


def test_safety_cap_not_completion(tmp_path):
    _,s=run(Client(),tmp_path,dict(CONFIG,max_steps=2))
    assert s["nav_activation_attempts"]==2 and s["termination_status"]=="INCOMPLETE_SAFETY_LIMIT"


def test_body_does_not_consume_budget(tmp_path):
    _,s=run(Client(),tmp_path)
    assert s["nav_activation_attempts"]==4 and s["candidate_policy"]=="GLOBAL_NAV_ONLY"


def test_no_selected_start_activates_all_five(tmp_path):
    _,s=run(Client(selected=None),tmp_path)
    assert s["nav_activation_attempts"]==5 and s["nav_items_verified"]==5


def test_success_ack_without_state_change_is_bounded(tmp_path):
    _,s=run(Client(change=False),tmp_path)
    assert s["destination_verification_failures"]==8 and s["nav_items_verified"]==1


def test_disabled_item_excluded():
    n=nodes();n[1]["enabled"]=False
    assert len(nav.discover(n,CONFIG))==4


def test_container_ancestry_discovery():
    n=nodes()
    for item in n[:5]:
        item.pop("role");item["viewIdResourceName"]="";item["ancestors"]=[dict(path="0.1",resource_id="app:id/bottom_navigation",class_name="HorizontalScrollView")]
    assert len(nav.discover(n,CONFIG))==5


def test_raw_xml_preserves_selected_and_ancestry():
    raw='<hierarchy><node resource-id="app:id/bottom_navigation"><node content-desc="Home" focusable="true" selected="true" bounds="[0,900][90,1000]"/></node></hierarchy>'
    items=nav.discover(nav.xml_nodes(raw),CONFIG)
    assert len(items)==1 and items[0]["selected"] and items[0]["evidence_source"]=="container_ancestry"


def test_state_visit_is_not_fabricated_physical_focus(tmp_path):
    rows,_=run(Client(),tmp_path)
    assert all(r["physical_visited"] is False and not r.get("merged_announcement") for r in rows)


def test_json_serialization(tmp_path):
    run(Client(),tmp_path)
    data=json.loads((tmp_path/"result.global_nav.json").read_text(encoding="utf-8"))
    assert data["summary"]["nav_items_verified"]==5 and len(data["transitions"])==4


def test_phase0a_termination_mapping():
    assert termination_status("global_nav_verified")=="COMPLETED"
    assert termination_status("global_nav_incomplete")=="INCOMPLETE_GLOBAL_NAV"
    assert termination_status("safety_limit")=="INCOMPLETE_SAFETY_LIMIT"


def test_collection_dispatch_keeps_content_path(monkeypatch,tmp_path):
    from tb_runner import collection_flow as flow
    calls=[]
    monkeypatch.setattr(flow,"_collect_tab_rows_inner",lambda *args,**kwargs:calls.append("content") or [])
    monkeypatch.setattr(flow,"_ensure_focusable_inventory",lambda *args:[])
    monkeypatch.setattr(flow,"_build_focusable_coverage_payload",lambda *args,**kwargs:dict(records=[]))
    monkeypatch.setattr(flow,"_save_focusable_coverage",lambda *args:None)
    for scenario in ("home_main","devices_main"):
        flow._collect_tab_rows_impl(SimpleNamespace(),"d",dict(scenario_id=scenario,scenario_type="content"),[],str(tmp_path/"out.xlsx"),str(tmp_path))
    assert calls==["content","content"]


def test_public_global_dispatch_bypasses_body_pipeline(monkeypatch,tmp_path):
    from tb_runner import collection_flow as flow
    monkeypatch.setattr(flow,"_collect_tab_rows_inner",lambda *args,**kwargs:pytest.fail("body pipeline entered"))
    monkeypatch.setattr(flow,"_ensure_focusable_inventory",lambda *args:[])
    monkeypatch.setattr(flow,"_build_focusable_coverage_payload",lambda *args,**kwargs:dict(records=[]))
    monkeypatch.setattr(flow,"_save_focusable_coverage",lambda *args:None)
    monkeypatch.setattr(flow,"save_excel_with_perf",lambda *args,**kwargs:None)
    c=Client();flow._collect_tab_rows_impl(c,"d",CONFIG,[],str(tmp_path/"out.xlsx"),str(tmp_path))
    summary=c.last_main_traversal_summary
    assert summary["nav_items_verified"]==5 and summary["traversal_complete"]
    assert summary["attempted_steps"]==summary["successful_moves"]==summary["unique_visited_instances"]==0


def test_evidence_serializes_all_destinations(tmp_path):
    events=[]
    c=Client();c.evidence_runtime=SimpleNamespace(is_enabled=True,start_scenario=lambda *a,**k:None,
                                                emit=lambda event,**kwargs:events.append((event,kwargs["payload"])))
    run(c,tmp_path)
    assert len([e for e in events if e[0]=="GLOBAL_NAV_TRANSITION"])==4
    assert next(p for e,p in events if e=="GLOBAL_NAV_SUMMARY")["nav_items_verified"]==5


def test_nav_excel_state_is_independent_of_speech_comparison(tmp_path):
    from tb_runner.excel_report import save_excel
    from openpyxl import load_workbook
    rows,_=run(Client(),tmp_path)
    save_excel(rows,str(tmp_path/"nav.xlsx"),with_images=False)
    book=load_workbook(tmp_path/"nav.xlsx",read_only=True,data_only=True)
    values=list(book["global_nav"].values)
    records=[dict(zip(values[0],r)) for r in values[1:]]
    assert len(records)==5 and all(r["nav_destination_verified"] for r in records)
    assert book["result"].max_row==1  # No fabricated speech comparison for state-only rows.


def test_optional_focus_error_does_not_erase_selected_verification(tmp_path):
    c=Client()
    def unavailable(**kwargs): raise RuntimeError("focus payload unavailable")
    c.get_focus=unavailable
    _,s=run(c,tmp_path)
    assert s["nav_items_verified"]==5


def test_after_dump_error_records_verification_failure(monkeypatch,tmp_path):
    original=nav.capture
    def fail_after(client,dev,config,folder,index):
        if index: raise RuntimeError("after dump unavailable")
        return original(client,dev,config,folder,index)
    monkeypatch.setattr(nav,"capture",fail_after)
    rows,s=run(Client(),tmp_path)
    assert s["destination_verification_failures"]==1 and s["termination_status"]=="INCOMPLETE_GLOBAL_NAV"
    assert rows[-1]["nav_transition"]["after_snapshot_error"]


def test_nav_metrics_do_not_increment_content_perf(tmp_path):
    from tb_runner.perf_stats import ScenarioPerfStats
    c=Client();perf=ScenarioPerfStats("global_nav_main","global nav")
    nav.collect(c,"d",CONFIG,[],str(tmp_path/"r.xlsx"),str(tmp_path),scenario_perf=perf)
    assert perf.total_steps==0 and perf.main_step_count==0
    assert c.last_main_traversal_summary["nav_recorded_rows"]==5


def test_local_role_tab_cannot_compete_with_global_container():
    n=nodes()
    for item in n[:5]:
        item["ancestors"]=[dict(path="nav",resource_id="app:id/bottom_navigation",class_name="BottomNavigationView")]
    n.append(dict(text="Home",role="tab",viewIdResourceName="local_tab",selected=True,focusable=True,boundsInScreen="0,100,100,200"))
    assert len(nav.discover(n,CONFIG))==5


def test_activation_ack_does_not_confirm_activation_without_state(tmp_path):
    rows,_=run(Client(change=False),tmp_path)
    assert rows[1]["nav_transition"]["activation_action_success"]
    assert not rows[1]["nav_transition"]["activation_confirmed"]


def test_marker_string_is_one_pattern():
    before,after=observation(None),observation(None)
    after["nodes"].append(dict(text="Devices root marker"))
    assert nav.verify_transition(before,after,before["items"][1],True,"Devices root marker")["destination_verified"]


def test_mutable_nav_badge_does_not_change_identity():
    a=nodes();b=nodes();b[4]["contentDescription"]="Menu, new content available"
    assert [i["instance_id"] for i in nav.discover(a,CONFIG)]==[i["instance_id"] for i in nav.discover(b,CONFIG)]
